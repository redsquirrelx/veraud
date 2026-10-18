from datetime import UTC, datetime
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from test_audit_diagnostics import incomplete_board
from test_audit_report import assessment, build, evidence, finding, populated

from server.api.audits import get_audit_service, router
from server.app.audit_service import AuditService, AuditState
from server.blackboard.impls.memory import MemoryBlackboard
from server.services.audit_report_docx import DOCX_MEDIA_TYPE, render_audit_report_docx


def document_text(document):
    return "\n".join([*[item.text for item in document.paragraphs],
        *[document_text(cell) for table in document.tables for row in table.rows for cell in row.cells]])


def test_same_json_produces_identical_valid_docx_with_all_ten_sections_and_traceability():
    board = populated()
    board.publish_finding(finding("discarded", "rejected", title="Hallazgo descartado"))
    report = build(board)
    first = render_audit_report_docx(report)
    second = render_audit_report_docx(report)
    assert first == second
    document = Document(BytesIO(first))
    headings = [item.text for item in document.paragraphs if item.style.name == "Heading 1"]
    assert len(headings) == 10
    assert [item.split(".")[0] for item in headings] == [str(index) for index in range(1, 11)]
    content = document_text(document)
    for expected in ("Seguridad", "No evaluado", "Descartado", "Hallazgo descartado", "Verificado",
            "domain/service.py",
            report.quality_results.findings[0].criterion.description):
        assert expected in content
    with ZipFile(BytesIO(first)) as archive:
        assert archive.testzip() is None
        assert all(item.date_time == (1980, 1, 1, 0, 0, 0) for item in archive.infolist())


def test_renderer_presents_source_as_text_and_replaces_invalid_xml_control_characters():
    board = populated()
    board.publish_finding(finding("control", description="Texto <script>literal</script>\x00 final"))
    report = build(board)
    output = render_audit_report_docx(report)
    assert "Texto <script>literal</script>\ufffd final" in document_text(Document(BytesIO(output)))


def test_nested_json_is_presented_as_word_tables_and_lists_with_generated_summary():
    report = build(incomplete_board())
    document = Document(BytesIO(render_audit_report_docx(report)))
    content = document_text(document)
    assert "HTTP 400" in content and "no constituyen defectos" in content
    assert "Ejecución X003" in content and "Llamadas al modelo" in content
    assert '"model_calls":' not in content and '"invocation_id":' not in content
    assert '"record_type":' not in content
    # Metrics and references must be actual Word table cells, not a JSON string.
    with ZipFile(BytesIO(render_audit_report_docx(report))) as archive:
        xml = archive.read("word/document.xml").decode()
        assert xml.count("<w:tbl>") > 15
        assert "w:tblHeader" in xml


def test_docx_references_are_short_internal_links_to_real_findings_and_evidence():
    finding_id = "6f2ba9b5-665b-4893-a3d6-f5840aef512c"
    evidence_id = "f9a0e5d4-2b41-457c-846d-c16525f0476b"
    board = MemoryBlackboard()
    board.publish_evidence(evidence(evidence_id, content="e1 = 'original source, not a reference'"))
    board.publish_finding(finding(finding_id, evidence_ids=(evidence_id,)))
    board.publish_assessment(assessment(finding_ids=(finding_id,), evidence_ids=(evidence_id,)))
    output = render_audit_report_docx(build(board))
    with ZipFile(BytesIO(output)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    name = lambda key: f"{{{ns['w']}}}{key}"
    visible = " ".join(node.text or "" for node in root.findall(".//w:t", ns))
    assert finding_id not in visible and evidence_id not in visible
    assert "Hallazgo H001" in visible and "Evidencia E001" in visible
    anchors = {node.get(name("name")) for node in root.findall(".//w:bookmarkStart", ns)}
    links = root.findall(".//w:hyperlink", ns)
    assert links and all(link.get(name("anchor")) in anchors for link in links)
    assert any(link.get(name("anchor")) == "ref_H001" for link in links)
    assert any(link.get(name("anchor")) == "ref_E001" for link in links)
    assert any(finding_id in link.get(name("tooltip"), "") for link in links)
    assert any(evidence_id in link.get(name("tooltip"), "") for link in links)


def test_tables_fit_the_page_and_do_not_contain_nested_tables():
    output = render_audit_report_docx(build(populated()))
    with ZipFile(BytesIO(output)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    name = lambda key: f"{{{ns['w']}}}{key}"
    assert not root.findall(".//w:tc/w:tbl", ns)
    for table in root.findall(".//w:tbl", ns):
        assert table.find("w:tblPr/w:tblLayout", ns).get(name("type")) == "fixed"
        width = int(table.find("w:tblPr/w:tblW", ns).get(name("w")))
        columns = sum(int(column.get(name("w"))) for column in table.findall("w:tblGrid/w:gridCol", ns))
        assert abs(columns - width) <= 2
        assert table.find("w:tr/w:trPr/w:tblHeader", ns) is not None


def test_source_content_is_not_rewritten_even_when_it_contains_record_ids():
    board = MemoryBlackboard()
    board.publish_evidence(evidence(content="f1 = e1  # literal source"))
    board.publish_finding(finding())
    output = render_audit_report_docx(build(board))
    with ZipFile(BytesIO(output)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert any(node.text == "f1 = e1  # literal source" for node in root.findall(".//w:t", ns))


def test_symbol_names_are_resolved_everywhere_with_real_destinations_and_literal_source():
    symbol_id, file_id = "symbol:" + "a" * 24, "file:" + "b" * 24
    missing = "component:" + "c" * 24
    board = MemoryBlackboard()
    board.publish_evidence(evidence(file_id=file_id, symbol_ids=(symbol_id,), content=f"# literal {symbol_id}", details={
        "investigation": {"plan": {"hypotheses": [{"id": "h1", "component_ids": [symbol_id]}]},
            "hypotheses": [{"id": "h1", "status": "confirmed", "evidence_ids": ["e1"]}]},
        "audit_context": {"entity_catalog": [
            {"id": file_id, "kind": "file", "path": "domain/service.py"},
            {"id": symbol_id, "kind": "symbol", "file_id": file_id, "name": "Service.parse_allowlist", "start_line": 10, "end_line": 25}]}}))
    board.publish_finding(finding(components=(symbol_id, missing), description=f"Inspect {symbol_id} in {file_id}."))
    report = build(board)
    assert report.annexes_and_traceability.reference_labels[symbol_id] == "Service.parse_allowlist — domain/service.py:10–25"
    assert any("Service.parse_allowlist — domain/service.py:10–25" in objective.text
        for objective in report.objectives_and_scope.specific_objectives)
    output = render_audit_report_docx(report)
    with ZipFile(BytesIO(output)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    visible = " ".join(node.text or "" for node in root.findall(".//w:t", ns))
    assert "Service.parse_allowlist — domain/service.py:10–25" in visible
    assert "Componente C001 — nombre no publicado" in visible
    assert missing not in visible and file_id not in visible
    assert visible.count(symbol_id) == 1  # Original source excerpt only.
    anchors = {node.get(f"{{{ns['w']}}}name") for node in root.findall(".//w:bookmarkStart", ns)}
    assert all(link.get(f"{{{ns['w']}}}anchor") in anchors for link in root.findall(".//w:hyperlink", ns))
    assert report.quality_results.findings[0].components == [symbol_id, missing]


def test_extensive_tables_and_annexes_have_collapsible_headings_without_hidden_rows():
    report = build(populated())
    evidence_item = report.annexes_and_traceability.evidence_inventory[0]
    inventory = [evidence_item.model_copy(update={"id": f"extra-evidence-{number}"}) for number in range(15)]
    report = report.model_copy(update={"annexes_and_traceability": report.annexes_and_traceability.model_copy(
        update={"evidence_inventory": inventory})})
    with ZipFile(BytesIO(render_audit_report_docx(report))) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "w15": "http://schemas.microsoft.com/office/word/2012/wordml"}
    folded = [
        p for p in root.findall(".//w:p", ns) if p.find("w:pPr/w15:collapsed", ns) is not None]
    assert any("Inventario de evidencias (15 registros)" in "".join(p.itertext()) for p in folded)
    assert all(p.find("w:pPr/w:pStyle", ns) is not None for p in folded)
    assert len(root.findall(".//w:bookmarkStart", ns)) >= 15
    assert not root.findall(".//w:vanish", ns)


def test_docx_endpoint_downloads_blackboard_report_and_returns_404_for_unknown_audit():
    service = AuditService(populated())
    service._states["audit"] = AuditState(run_id="audit", invocation_id="manager", status="completed",
        view_url="/blackboard?run_id=audit", created_at=datetime.now(UTC))
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_audit_service] = lambda: service
    with TestClient(app) as client:
        response = client.get("/api/audits/audit/report.docx")
        assert response.status_code == 200
        assert response.headers["content-type"] == DOCX_MEDIA_TYPE
        assert response.headers["content-disposition"] == 'attachment; filename="informe-auditoria.docx"'
        assert response.headers["cache-control"] == "no-store"
        assert "Auditar mantenibilidad." in document_text(Document(BytesIO(response.content)))
        assert client.get("/api/audits/unknown/report.docx").status_code == 404
        # JSON remains the canonical representation and has its original endpoint.
        assert client.get("/api/audits/audit/report").json()["schema_version"] == "1.0"
